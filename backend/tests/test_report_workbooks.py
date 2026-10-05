"""Phase 14: hand-worked six-section reports, reopened and downloaded.

The independent controls use synthetic identities and prices. They reproduce
credits, transferred accounts and incomplete history from the private source
review without publishing business data or treating a cached workbook as truth.
"""

from __future__ import annotations

import io
import zipfile
from datetime import date, datetime

import openpyxl
import polars as pl
import pytest

from app.actions.base import ActionResult
from app.actions.monthly_sales_rep_report import MonthlySalesRepReportAction
from app.errors import ActionExecutionError
from app.models.artifact import Artifact, check_artifact_filename
from app.models.report_spec import REP_COLUMN, TRANSACTION_COLUMNS
from app.models.run import RunResult
from app.services import report_workbooks
from app.services.data_library import ensure_known_datasets
from app.services.ingestion import (
    SourceFile, commit_account_assignments, commit_monthly_sales,
    commit_monthly_samples,
)
from app.services.monthly_report import build_tables, prepare
from app.services.report_views import (
    BOTTLES, CHANGE, COMPANY_SALES, COMPANY_SHARE, CURRENT_R12, GROWTH,
    NET_SALES, PRIOR_R12, PRODUCT, REP_SALES, REP_SHARE, R12_BOTTLES, VIEW_IDS,
)
from app.services.report_workbooks import render_rep_workbooks, report_sheets, workbook_names
from app.services.runner import execute_run
from app.services.workbook import Sheet, render_workbook
from tests.fixtures import completed_report_month as source

SHEET_NAMES = [
    "Samples August 2026", "Samples R12", "Sales R12 by Account",
    "Sales August 2026", "Sales by Product and Account", "Sales by Account R12",
]
SUPPLIER = TRANSACTION_COLUMNS.supplier
CUSTOMER = TRANSACTION_COLUMNS.customer
DATE = TRANSACTION_COLUMNS.date


@pytest.fixture
def completed():
    inputs = source.inputs()
    prepared = prepare(inputs["sales_history"], inputs["sample_history"], inputs["account_assignments"])
    assert prepared.usable, prepared.errors
    return prepared, build_tables(prepared)


def only(frame: pl.DataFrame, **keys):
    for key, value in keys.items():
        frame = frame.filter(pl.col(key) == value)
    assert frame.height == 1
    return frame.row(0, named=True)


def footer(tables, rep, section):
    return dict(tables["workbook_totals"].filter(
        (pl.col(REP_COLUMN) == rep) & (pl.col("Section") == section)
    ).select("Measure", "Value").iter_rows())


def test_transfers_keep_both_invoice_reps_without_snapshot_only_reps(completed):
    prepared, tables = completed
    assert prepared.reps == (source.ALPHA, source.BETA)
    assert "UNRECOGNISED_SALES_REP" not in [issue.code for issue in prepared.warnings]
    comparison = tables["rolling_account_comparison"]
    alpha = only(comparison, **{REP_COLUMN: source.ALPHA, CUSTOMER: source.TRANSFERRED})
    beta = only(comparison, **{REP_COLUMN: source.BETA, CUSTOMER: source.TRANSFERRED})
    assert (alpha[PRIOR_R12], alpha[CURRENT_R12], alpha[CHANGE], alpha[GROWTH], alpha["Status"]) == (0, 42.75, 42.75, None, "New")
    assert (beta[PRIOR_R12], beta[CURRENT_R12], beta["Status"]) == (0, 14.25, "New")
    assert only(tables["company_summary"])["Revenue"] == 58.25
    assert tables["rep_summary"].filter(pl.col(REP_COLUMN) == source.IDLE).height == 0


def test_sample_credits_and_twelve_calendar_months_have_exact_values(completed):
    _, tables = completed
    monthly = only(tables["monthly_samples"], **{REP_COLUMN: source.ALPHA})
    assert monthly[BOTTLES] == 0  # August sample and sample credit cancel.
    assert only(tables["monthly_samples"], **{REP_COLUMN: source.BETA})[BOTTLES] == 2
    rolling = only(tables["rolling_samples"], **{REP_COLUMN: source.ALPHA})
    assert list(rolling)[3:15] == [
        "Sep 2025", "Oct 2025", "Nov 2025", "Dec 2025", "Jan 2026", "Feb 2026",
        "Mar 2026", "Apr 2026", "May 2026", "Jun 2026", "Jul 2026", "Aug 2026",
    ]
    assert [rolling[month] for month in list(rolling)[3:15]] == [1] * 11 + [0]
    assert rolling[R12_BOTTLES] == 11
    assert footer(tables, source.ALPHA, "rolling_samples")[R12_BOTTLES] == 11


def test_r12_account_and_product_controls_are_independently_worked(completed):
    _, tables = completed
    sales = tables["rolling_account_sales"]
    assert only(sales, **{REP_COLUMN: source.ALPHA, CUSTOMER: source.ACCOUNT})[NET_SALES] == 15
    assert only(sales, **{REP_COLUMN: source.ALPHA, CUSTOMER: source.TRANSFERRED})[NET_SALES] == 42.75
    assert footer(tables, source.ALPHA, "rolling_account_sales")[NET_SALES] == 57.75
    products = tables["rolling_product_accounts"]
    assert only(products, **{REP_COLUMN: source.ALPHA, CUSTOMER: source.ACCOUNT})[BOTTLES] == 12
    assert only(products, **{REP_COLUMN: source.ALPHA, CUSTOMER: source.TRANSFERRED})[BOTTLES] == 9
    assert only(products, **{REP_COLUMN: source.BETA, CUSTOMER: source.TRANSFERRED})[BOTTLES] == 3
    totals = footer(tables, source.ALPHA, "rolling_account_comparison")
    assert (totals[PRIOR_R12], totals[CURRENT_R12], totals[CHANGE], totals[GROWTH]) == (15, 57.75, 42.75, 2.85)


def test_supplier_percentages_and_footer_are_ratios_of_correct_amounts(completed):
    _, tables = completed
    suppliers = tables["monthly_supplier_sales"]
    alpha = only(suppliers, **{REP_COLUMN: source.ALPHA, SUPPLIER: source.SUPPLIER_TWO})
    assert alpha[REP_SALES] == 42.75
    assert alpha[COMPANY_SALES] == 57
    assert alpha[REP_SHARE] == pytest.approx(42.75 / 44)
    assert alpha[COMPANY_SHARE] == 0.75
    beta = only(suppliers, **{REP_COLUMN: source.BETA})
    assert (beta[REP_SALES], beta[REP_SHARE], beta[COMPANY_SALES], beta[COMPANY_SHARE]) == (14.25, 1, 57, 0.25)
    alpha_footer = footer(tables, source.ALPHA, "monthly_supplier_sales")
    assert alpha_footer[COMPANY_SHARE] == pytest.approx(44 / 58.25)
    assert alpha_footer[COMPANY_SHARE] != pytest.approx(suppliers.filter(pl.col(REP_COLUMN) == source.ALPHA)[COMPANY_SHARE].sum())
    # Beta displays only Supplier Two; the footer covers the supplier shown.
    assert footer(tables, source.BETA, "monthly_supplier_sales")[COMPANY_SALES] == 57


@pytest.mark.parametrize("rep", [source.ALPHA, source.BETA, source.IDLE])
def test_all_six_sheets_round_trip_every_data_and_total_value(completed, rep):
    prepared, tables = completed
    sheets = report_sheets(prepared, tables, rep)
    workbook = openpyxl.load_workbook(io.BytesIO(render_workbook(sheets)))
    assert workbook.sheetnames == SHEET_NAMES
    assert all(sheet.sheet_state == "visible" for sheet in workbook)
    for descriptor, sheet, view_id in zip(sheets, workbook, VIEW_IDS, strict=True):
        expected = tables[view_id].filter(pl.col(REP_COLUMN) == rep).drop(REP_COLUMN)
        # Header position includes title, period, notes and a spacer.
        header_row = 4 + len(descriptor.notes)
        assert [cell.value for cell in sheet[header_row]] == expected.columns
        assert sheet.freeze_panes == f"A{header_row + 1}"
        assert sheet.sheet_view.showGridLines is False
        assert sheet.row_dimensions[header_row].height == 34
        assert len(sheet.tables) == 1
        table = next(iter(sheet.tables.values()))
        assert table.autoFilter is not None
        assert table.tableStyleInfo.name == "TableStyleMedium2"
        assert all(sheet.column_dimensions[letter].width > 0 for letter in ["A", "B"])
        for offset, row in enumerate(expected.iter_rows(), start=header_row + 1):
            for column, value in enumerate(row, start=1):
                cell = sheet.cell(offset, column)
                if isinstance(value, float):
                    assert cell.value == pytest.approx(value, abs=1e-9)
                    assert cell.data_type == "n"
                else:
                    assert cell.value == value
                if expected.columns[column - 1] in (REP_SHARE, COMPANY_SHARE, GROWTH) and value is not None:
                    assert "%" in cell.number_format and cell.data_type == "n"
                if expected.columns[column - 1] in (NET_SALES, REP_SALES, COMPANY_SALES, PRIOR_R12, CURRENT_R12, CHANGE) and value is not None:
                    assert "$" in cell.number_format and cell.data_type == "n"
        totals = footer(tables, rep, view_id)
        total_row = header_row + expected.height + 1
        assert sheet.cell(total_row, 1).value == "Total"
        for column, name in enumerate(expected.columns, start=1):
            if name in totals:
                value = sheet.cell(total_row, column).value
                if totals[name] is None:
                    assert value is None
                else:
                    assert value == pytest.approx(totals[name])
        assert all(cell.data_type != "f" for row in sheet for cell in row)


def test_partial_r12_history_has_blank_totals_and_explicit_notes():
    inputs = source.inputs()
    inputs["sales_history"] = inputs["sales_history"].filter(~pl.col(DATE).str.starts_with("2026-03"))
    inputs["sample_history"] = inputs["sample_history"].filter(~pl.col(DATE).str.starts_with("2026-04"))
    prepared = prepare(inputs["sales_history"], inputs["sample_history"], inputs["account_assignments"])
    assert prepared.usable
    tables = build_tables(prepared)
    assert tables["rolling_account_sales"][NET_SALES].null_count() == tables["rolling_account_sales"].height
    assert tables["rolling_product_accounts"][BOTTLES].null_count() == tables["rolling_product_accounts"].height
    assert tables["rolling_samples"]["Apr 2026"].null_count() == tables["rolling_samples"].height
    assert footer(tables, source.ALPHA, "rolling_samples")[R12_BOTTLES] is None
    comparison = only(tables["rolling_account_comparison"], **{REP_COLUMN: source.ALPHA, CUSTOMER: source.ACCOUNT})
    assert comparison[PRIOR_R12] == 15
    assert (comparison[CURRENT_R12], comparison[CHANGE], comparison[GROWTH], comparison["Status"]) == (None, None, None, "Incomplete history")
    # The complete reporting month is still usable on its own.
    assert footer(tables, source.ALPHA, "monthly_supplier_sales")[REP_SALES] == 44
    sheets = report_sheets(prepared, tables, source.ALPHA)
    assert any("2026-04" in note for note in sheets[1].notes)
    assert any("2026-03" in note for note in sheets[5].notes)


def test_samples_need_an_invoice_rep_but_not_an_account():
    inputs = source.inputs()
    inputs["sample_history"] = inputs["sample_history"].with_columns(pl.lit(None, pl.String).alias(CUSTOMER))
    action = MonthlySalesRepReportAction()
    assert action.validate(inputs) == []
    assert len(action.run(inputs).artifacts) == 2


def test_duplicate_sales_and_samples_in_earlier_r12_months_are_reported_and_retained():
    inputs = source.inputs()
    inputs["sales_history"] = inputs["sales_history"].vstack(inputs["sales_history"].filter(pl.col(DATE).str.starts_with("2026-02")))
    inputs["sample_history"] = inputs["sample_history"].vstack(inputs["sample_history"].filter(pl.col(DATE).str.starts_with("2026-03")))
    prepared = prepare(inputs["sales_history"], inputs["sample_history"], inputs["account_assignments"], sales_slot="sales_history", samples_slot="sample_history")
    repeated = [issue for issue in prepared.warnings if issue.code == "DUPLICATE_SOURCE_ROWS"]
    assert {issue.slot_id for issue in repeated} == {"sales_history", "sample_history"}
    assert all(issue.details["row_count"] == 1 for issue in repeated)
    tables = build_tables(prepared)
    assert footer(tables, source.ALPHA, "rolling_account_sales")[NET_SALES] == 59
    assert footer(tables, source.ALPHA, "rolling_samples")[R12_BOTTLES] == 12


@pytest.mark.parametrize("kind", ["Sample Invoice", "Rebate", "", None])
def test_wrong_or_unknown_sales_type_fails_before_rendering(kind):
    inputs = source.inputs()
    inputs["sales_history"] = inputs["sales_history"].with_columns(pl.lit(kind, pl.String).alias(TRANSACTION_COLUMNS.invoice_type))
    assert "UNEXPECTED_INVOICE_TYPE" in [issue.code for issue in MonthlySalesRepReportAction().validate(inputs)]


def test_known_snapshot_totals_are_preserved_but_never_used_for_performance():
    inputs = source.inputs()
    inputs["account_assignments"] = inputs["account_assignments"].with_columns(
        pl.lit(987654).alias("Prior R12"), pl.lit(987654).alias("Current R12"),
        pl.lit(987654).alias("$ Change"), pl.lit(987654).alias("% Change"),
    )
    prepared = prepare(inputs["sales_history"], inputs["sample_history"], inputs["account_assignments"])
    assert "UNEXPECTED_SOURCE_COLUMNS" not in [issue.code for issue in prepared.warnings]
    assert footer(build_tables(prepared), source.ALPHA, "rolling_account_comparison")[CURRENT_R12] == 57.75
    assert inputs["account_assignments"]["Current R12"].sum() == 3 * 987654


def test_workbook_names_are_flat_case_insensitive_and_preserve_period():
    names = workbook_names(["Rep/A", "Rep:A", "rep-a", "é" * 250, "😀" * 250], "September 2026")
    assert len(set(name.casefold() for name in names)) == len(names)
    assert all("September 2026" in name and name.endswith(".xlsx") for name in names)
    assert all(check_artifact_filename(name) == name for name in names)


def test_one_render_policy_handles_calendar_labels_across_year_boundary():
    inputs = source.inputs()
    inputs["sales_history"] = inputs["sales_history"].filter(pl.col(DATE) <= "2026-01-31")
    inputs["sample_history"] = inputs["sample_history"].filter(pl.col(DATE) <= "2026-01-31")
    prepared = prepare(inputs["sales_history"], inputs["sample_history"], inputs["account_assignments"])
    tables = build_tables(prepared)
    sheets = report_sheets(prepared, tables, source.ALPHA)
    assert sheets[0].name == "Samples January 2026"
    assert sheets[3].name == "Sales January 2026"
    assert tables["rolling_samples"].columns[3:15] == [
        "Feb 2025", "Mar 2025", "Apr 2025", "May 2025", "Jun 2025", "Jul 2025",
        "Aug 2025", "Sep 2025", "Oct 2025", "Nov 2025", "Dec 2025", "Jan 2026",
    ]


def test_workbooks_are_byte_reproducible_and_keep_input_frames(completed):
    inputs = source.inputs()
    originals = {key: frame.clone() for key, frame in inputs.items()}
    first = MonthlySalesRepReportAction().run(inputs)
    second = MonthlySalesRepReportAction().run({key: frame.reverse() for key, frame in inputs.items()})
    assert [(a.id, a.filename, a.payload) for a in first.artifacts] == [(a.id, a.filename, a.payload) for a in second.artifacts]
    for key, frame in inputs.items():
        assert frame.equals(originals[key])
    workbook = openpyxl.load_workbook(io.BytesIO(first.artifacts[0].payload))
    assert workbook.properties.created == datetime(2000, 1, 1)


def test_user_text_in_titles_notes_and_tables_stays_literal():
    payload = render_workbook([Sheet(name="Safe", frame=pl.DataFrame({"Account": ["=1+1", "https://example.test"]}),
        title="=HYPERLINK(\"https://example.test\")", notes=("=1+1",), total_row={})])
    workbook = openpyxl.load_workbook(io.BytesIO(payload))
    assert all(cell.data_type != "f" for row in workbook["Safe"] for cell in row)
    assert workbook["Safe"]["A1"].value == '=HYPERLINK("https://example.test")'
    assert workbook["Safe"]["A2"].value == "=1+1"


@pytest.mark.parametrize("filename", ["../reports.zip", "reports.xlsx", "reports\n.zip"])
def test_bundle_name_is_validated_at_both_runtime_boundaries(filename):
    artifact = Artifact.workbook(id="one", label="One", filename="one.xlsx", payload=b"bytes")
    tables = {"result": pl.DataFrame({"n": [1]})}
    with pytest.raises(ValueError):
        ActionResult(outputs=tables, artifacts=(artifact,), artifact_bundle_filename=filename)
    with pytest.raises(ValueError):
        RunResult.of(tables, {artifact.id: artifact}, filename)


@pytest.fixture
def committed_report(data_library):
    ensure_known_datasets()
    inputs = source.inputs()
    for slot, commit in (("sales_history", commit_monthly_sales), ("sample_history", commit_monthly_samples)):
        frame = inputs[slot]
        for month in sorted(frame[DATE].str.slice(0, 7).unique().to_list()):
            rows = frame.filter(pl.col(DATE).str.starts_with(month))
            commit(SourceFile(filename=f"{slot}-{month}.csv", payload=rows.write_csv().encode()), expected_period=month, today=date(2026, 10, 1))
    commit_account_assignments(SourceFile(filename="assignments.csv", payload=inputs["account_assignments"].write_csv().encode()), period="2026-08")
    return {"sales_history": "history:2026-08", "sample_history": "history:2026-08", "account_assignments": "period:2026-08"}


def test_one_http_run_downloads_every_workbook_individually_and_as_month_zip(committed_report, client, run_store, quarantine):
    response = client.post("/api/runs", data={"action_id": "monthly_sales_rep_report", **committed_report})
    assert response.status_code == 200, response.text
    manifest = response.json()
    run = run_store.get_run(manifest["run_id"])
    assert run.result is not None
    assert len(manifest["artifacts"]) == 2
    bundle = client.get(f"/api/runs/{run.run_id}/artifacts/download/zip")
    assert bundle.status_code == 200
    assert 'filename="August 2026 Sales Rep Reports.zip"' in bundle.headers["content-disposition"]
    with zipfile.ZipFile(io.BytesIO(bundle.content)) as archive:
        assert archive.namelist() == [f"{rep} - August 2026.xlsx" for rep in (source.ALPHA, source.BETA)]
        for artifact in manifest["artifacts"]:
            individual = client.get(f"/api/runs/{run.run_id}/artifacts/{artifact['id']}/download")
            assert individual.status_code == 200
            assert individual.content == archive.read(artifact["filename"])
            workbook = openpyxl.load_workbook(io.BytesIO(individual.content))
            assert workbook.sheetnames == SHEET_NAMES
    assert list(quarantine.iterdir()) == []


def test_a_render_failure_does_not_publish_a_partial_batch(committed_report, monkeypatch, run_store):
    actual = report_workbooks.render_workbook
    rendered = []

    def fail_on_second(sheets):
        rendered.append(sheets)
        if len(rendered) == 2:
            raise ValueError("Synthetic second-workbook failure")
        return actual(sheets)

    monkeypatch.setattr(report_workbooks, "render_workbook", fail_on_second)
    with pytest.raises(ActionExecutionError):
        execute_run(MonthlySalesRepReportAction(), uploads={}, dataset_references=committed_report)
    run = run_store.list_runs()[0]
    assert len(rendered) == 2
    assert run.status.value == "failed"
    assert run.result is None
    assert run.artifacts == ()
