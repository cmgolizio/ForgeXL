"""Phase 15: exercise browser requests, persistence, replay, and real downloads.

Expected sales values come from the hand-worked golden fixture. All files here
are synthetic; rendering and source selection use the registered report Action.
"""
from __future__ import annotations

from datetime import timedelta
from io import BytesIO
import zipfile

from openpyxl import load_workbook
import pytest

from app.actions import monthly_sales_rep_report as report_action
from app.errors import DataLibraryError
from app.models.library import now, new_version_id
from app.services import cycle_receipts, data_library as library, history_workflow, monthly_workflow, run_store
from app.services.data_library import LocalDataLibrary
from app.services.run_store import InMemoryRunStore
from tests.fixtures import report_months as golden
from tests.fixtures.monthly_sources import assignments, transaction_row, transactions

PERIOD = golden.GOLDEN_MONTH
PREFIX = "/api/monthly"


def source(rows):
    return transactions(rows).as_csv()


def monthly_files(*, sales=None, samples=None, owners=None):
    return {
        "sales_history": ("sales.csv", sales if sales is not None else source(golden.SALES_ROWS[PERIOD])),
        "sample_history": ("samples.csv", samples if samples is not None else source(golden.SAMPLE_ROWS[PERIOD])),
    }


def history(client, dataset_id, rows):
    response = client.post(PREFIX + "/history/validate", data={"dataset_id": dataset_id},
        files={"source_file": ("history.csv", source(rows))})
    assert response.status_code == 200, response.text
    review = response.json()
    assert review["ready"], review
    saved = client.post(PREFIX + "/history/commit", json={"validation_id": review["validation_id"], "acknowledge_warnings": True})
    assert saved.status_code == 200 and saved.json()["status"] == "saved", saved.text
    return review, saved.json()


def seed_history(client):
    history(client, "sales_history", [row for month, rows in golden.SALES_ROWS.items() if month != PERIOD for row in rows])
    history(client, "sample_history", golden.SAMPLE_ROWS["2026-08"])


def validate(client, files=None, **data):
    response = client.post(PREFIX + "/validate", data={"period": PERIOD, **data}, files=files or {})
    assert response.status_code == 200, response.text
    return response.json()


def generate(client, review):
    response = client.post(PREFIX + "/generate", json={"validation_id": review["validation_id"], "acknowledge_warnings": True})
    assert response.status_code == 200, response.text
    return response.json()


def cycle(client):
    seed_history(client)
    review = validate(client, monthly_files())
    assert review["ready"], review
    result = generate(client, review)
    assert result["status"] == "reports_generated", result
    return review, result


def preview(client, result, table="company_summary"):
    response = client.get(f'/api/runs/{result["manifest"]["run_id"]}/outputs/{table}/preview')
    assert response.status_code == 200
    payload = response.json()
    return [dict(zip(payload["columns"], row, strict=True)) for row in payload["rows"]]


def archive(client, result):
    response = client.get(f'/api/runs/{result["manifest"]["run_id"]}/artifacts/download/zip')
    assert response.status_code == 200
    assert "September 2026" in response.headers["content-disposition"]
    with zipfile.ZipFile(BytesIO(response.content)) as bundle:
        return {name: bundle.read(name) for name in bundle.namelist()}


def test_validate_then_generate_commits_once_and_downloads_every_rep(client):
    seed_history(client)
    review = validate(client, monthly_files())
    assert review["ready"] and review["warnings"]
    assert set(review["reps"]) == set(golden.EXPECTED["reps"])
    # A successful upload review has not committed the new month.
    assert all(v.period != PERIOD for v in library.current_versions("sales_history"))
    assert any(item["missing_months"] for item in review["coverage"])
    refused = client.post(PREFIX + "/generate", json={"validation_id": review["validation_id"]})
    assert refused.status_code == 400
    result = generate(client, review)
    assert result["sources_committed"] and len(result["committed_versions"]) == 2
    assert preview(client, result)[0]["Revenue"] == 995.0
    assert preview(client, result)[0]["Quantity"] == 33.0
    assert {item["code"] for item in result["manifest"]["validation"]["warnings"]} >= {item["code"] for item in review["warnings"]}
    files = archive(client, result)
    assert len(files) == 2
    for name, payload in files.items():
        assert name.endswith(" - September 2026.xlsx")
        workbook = load_workbook(BytesIO(payload), data_only=True)
        assert len(workbook.sheetnames) == 6
        assert all(sheet.max_row > 1 for sheet in workbook)
        workbook.close()
    assert client.post(PREFIX + "/generate", json={"validation_id": review["validation_id"], "acknowledge_warnings": True}).status_code == 400


def test_restart_replays_exact_versions_with_identical_workbooks(client, data_library, monkeypatch):
    _, original = cycle(client)
    first_zip = archive(client, original)
    first_summary = preview(client, original)
    monkeypatch.setattr(library, "DATA_LIBRARY", LocalDataLibrary(data_library.root))
    monkeypatch.setattr(cycle_receipts, "CYCLE_RECEIPTS", cycle_receipts.CycleReceiptStore(data_library.root / ".reporting-cycles"))
    monkeypatch.setattr(monthly_workflow, "WORKFLOW", monthly_workflow.MonthlyWorkflow())
    monkeypatch.setattr(history_workflow, "HISTORY_WORKFLOW", history_workflow.HistoryWorkflow())
    monkeypatch.setattr(run_store, "RUN_STORE", InMemoryRunStore())
    assert client.get(f'/api/runs/{original["manifest"]["run_id"]}').status_code == 404
    catalog = client.get(PREFIX + "/catalog").json()
    assert next(p for p in catalog["periods"] if p["period"] == PERIOD)["cycles"][0]["cycle_id"] == original["receipt"]["cycle_id"]
    review = client.post(PREFIX + "/validate-saved", json={"period": PERIOD, "cycle_id": original["receipt"]["cycle_id"]}).json()
    assert review["ready"]
    repeated = generate(client, review)
    assert repeated["receipt"] == original["receipt"]
    assert repeated["committed_versions"] == {}
    assert preview(client, repeated) == first_summary
    assert archive(client, repeated) == first_zip


def test_duplicate_and_correction_do_not_append_and_old_cycle_keeps_original(client):
    _, original = cycle(client)
    old_id = original["committed_versions"]["sales_history"]
    before = len(library.list_versions("sales_history"))
    duplicate = validate(client, {"sales_history": monthly_files()["sales_history"]})
    assert duplicate["ready"]
    assert len(library.list_versions("sales_history")) == before
    corrected = source([*golden.SALES_ROWS[PERIOD], transaction_row(invoice_date="2026-09-22", total_price=100, quantity=2)])
    refused = validate(client, {"sales_history": ("corrected.csv", corrected)})
    assert not refused["ready"]
    no_reason = validate(client, {"sales_history": ("corrected.csv", corrected)}, **{"sales_history.replaces": old_id})
    assert "CORRECTION_REASON_REQUIRED" in {issue["code"] for issue in no_reason["errors"]}
    review = validate(client, {"sales_history": ("corrected.csv", corrected)}, reason="Missing invoice included", **{"sales_history.replaces": old_id})
    assert review["ready"], review
    result = generate(client, review)
    assert preview(client, result)[0]["Revenue"] == 1095.0
    assert len(library.list_versions("sales_history")) == before + 1
    current = library.current_version("sales_history", PERIOD)
    assert current.supersedes == old_id
    assert library.get_version("sales_history", old_id).row_count == 5
    replay = client.post(PREFIX + "/validate-saved", json={"period": PERIOD, "cycle_id": original["receipt"]["cycle_id"]}).json()
    assert preview(client, generate(client, replay))[0]["Revenue"] == 995.0
    fresh = client.post(PREFIX + "/validate-saved", json={"period": PERIOD}).json()
    assert "Current stored versions" in fresh["source_selection"]
    captured = generate(client, fresh)
    assert captured["receipt"]["versions"]["sales_history"][-1] == current.version_id
    assert preview(client, captured)[0]["Revenue"] == 1095.0


def test_assignment_upload_is_rejected_instead_of_being_optional(client):
    response = client.post(PREFIX + "/validate", data={"period": PERIOD},
        files={"account_assignments": ("owners.csv", golden.assignment_table().as_csv())})
    assert response.status_code == 400
    assert library.list_datasets() == []

def test_render_failure_keeps_receipt_and_sources_for_retry(client, monkeypatch):
    seed_history(client)
    review = validate(client, monthly_files())
    render = report_action.render_rep_workbooks
    def broken(*args, **kwargs):
        raise RuntimeError("private /Users/secret/workbook.xlsx")
    monkeypatch.setattr(report_action, "render_rep_workbooks", broken)
    failed = generate(client, review)
    assert failed["status"] == "generation_failed" and failed["sources_committed"]
    assert failed["receipt"] and not failed["manifest"]
    assert "secret" not in str(failed)
    monkeypatch.setattr(report_action, "render_rep_workbooks", render)
    repeated_review = client.post(PREFIX + "/validate-saved", json={"period": PERIOD, "cycle_id": failed["receipt"]["cycle_id"]}).json()
    result = generate(client, repeated_review)
    assert result["status"] == "reports_generated"
    assert result["committed_versions"] == {}
    assert preview(client, result)[0]["Revenue"] == 995.0


def test_partial_commit_is_reported_and_missing_sources_can_be_resubmitted(client, monkeypatch):
    seed_history(client)
    review = validate(client, monthly_files())
    commit = monthly_workflow.ingestion.commit_validated_source
    def fail_samples(check, **kwargs):
        if check.dataset_id == "sample_history":
            raise DataLibraryError("Synthetic disk failure")
        return commit(check, **kwargs)
    monkeypatch.setattr(monthly_workflow.ingestion, "commit_validated_source", fail_samples)
    result = generate(client, review)
    assert result["status"] == "commit_failed" and not result["sources_committed"]
    assert set(result["committed_versions"]) == {"sales_history"}
    assert result["receipt"] is None
    monkeypatch.setattr(monthly_workflow.ingestion, "commit_validated_source", commit)
    remaining = monthly_files(); remaining.pop("sales_history")
    repeated = generate(client, validate(client, remaining))
    assert repeated["status"] == "reports_generated"
    assert len([v for v in library.list_versions("sales_history") if v.period == PERIOD]) == 1


def test_receipt_write_failure_keeps_sources_and_allows_current_source_retry(client, monkeypatch):
    seed_history(client)
    review = validate(client, monthly_files())
    with monkeypatch.context() as patch:
        patch.setattr(cycle_receipts.CYCLE_RECEIPTS, "save", lambda record: (_ for _ in ()).throw(DataLibraryError("Receipt unavailable")))
        failed = generate(client, review)
    assert failed["status"] == "commit_failed" and failed["sources_committed"]
    retry = client.post(PREFIX + "/validate-saved", json={"period": PERIOD}).json()
    assert generate(client, retry)["status"] == "reports_generated"


@pytest.mark.parametrize("fault", ["period", "missing", "number", "same"])
def test_invalid_inputs_never_commit_any_month(client, fault):
    files = monthly_files()
    if fault == "period": files["sales_history"] = ("wrong.csv", source(golden.SALES_ROWS["2026-08"]))
    if fault == "missing": files.pop("sample_history")
    if fault == "number": files["sales_history"] = ("bad.csv", source([transaction_row(invoice_date="2026-09-01", total_price="not-a-number")]))
    if fault == "same": files["sample_history"] = files["sales_history"]
    review = validate(client, files)
    assert not review["ready"] and not review["validation_id"] and review["errors"]
    assert library.list_datasets() == []


def test_expired_review_and_stale_correction_are_rejected(client, monkeypatch):
    _, original = cycle(client)
    review = client.post(PREFIX + "/validate-saved", json={"period": PERIOD}).json()
    monkeypatch.setattr(monthly_workflow, "now", lambda: now() + timedelta(hours=1))
    assert client.post(PREFIX + "/generate", json={"validation_id": review["validation_id"], "acknowledge_warnings": True}).status_code == 400
    stale = validate(client, {"sales_history": ("changed.csv", source([transaction_row(invoice_date="2026-09-01")]))}, reason="Correction", **{"sales_history.replaces": new_version_id()})
    assert not stale["ready"]
    assert "STALE_CORRECTION" in {issue["code"] for issue in stale["errors"]}


def test_changed_live_sources_invalidate_review_but_exact_cycle_survives(client):
    _, original = cycle(client)
    live_review = client.post(PREFIX + "/validate-saved", json={"period": PERIOD}).json()
    from app.services.ingestion import SourceFile, commit_monthly_sales
    commit_monthly_sales(SourceFile("changed.csv", source([transaction_row(invoice_date="2026-09-01")])),
        expected_period=PERIOD, replaces=original["committed_versions"]["sales_history"], reason="Concurrent correction")
    assert client.post(PREFIX + "/generate", json={"validation_id": live_review["validation_id"], "acknowledge_warnings": True}).status_code == 400
    exact = client.post(PREFIX + "/validate-saved", json={"period": PERIOD, "cycle_id": original["receipt"]["cycle_id"]}).json()
    assert preview(client, generate(client, exact))[0]["Revenue"] == 995.0


def test_history_bootstrap_refuses_duplicate_months_and_accepts_a_missing_month(client):
    rows = [*golden.SALES_ROWS["2025-08"], *golden.SALES_ROWS["2025-09"]]
    review, result = history(client, "sales_history", rows)
    assert review["operation"] == "bootstrap" and len(result["committed_versions"]) == 2
    assert [(v.period, v.row_count) for v in library.current_versions("sales_history")] == [("2025-08", 2), ("2025-09", 3)]
    repeated = client.post(PREFIX + "/history/validate", data={"dataset_id": "sales_history"}, files={"source_file": ("history.csv", source(rows))}).json()
    assert not repeated["ready"]
    history(client, "sales_history", golden.SALES_ROWS["2026-08"])
    assert len(library.current_versions("sales_history")) == 3


@pytest.mark.parametrize("period", ["0000-01", "0001-01", "2026-13", "../../private"])
def test_period_path_boundary_is_structured(client, period):
    response = client.post(PREFIX + "/validate-saved", json={"period": period})
    assert response.status_code == 400 and response.json()["error"]["code"] == "INVALID_REQUEST"


def test_catalog_contains_metadata_only_and_no_local_paths(client):
    cycle(client)
    payload = client.get(PREFIX + "/catalog").json()
    assert len(payload["datasets"]) == 2
    assert "/workspace/" not in str(payload) and "parquet" not in str(payload)
    assert "rows" not in payload["datasets"][0]


def test_an_obsolete_token_does_not_destroy_the_new_review(client):
    first = validate(client, monthly_files())
    second = validate(client, monthly_files())
    refused = client.post(PREFIX + "/generate", json={"validation_id": first["validation_id"], "acknowledge_warnings": True})
    assert refused.status_code == 400
    assert generate(client, second)["status"] == "reports_generated"


def test_discard_releases_the_upload_review(client):
    review = validate(client, monthly_files())
    assert monthly_workflow.WORKFLOW._pending is not None
    assert client.post(PREFIX + "/discard", json={"validation_id": review["validation_id"]}).status_code == 200
    assert monthly_workflow.WORKFLOW._pending is None
    assert client.post(PREFIX + "/generate", json={"validation_id": review["validation_id"]}).status_code == 400


@pytest.mark.parametrize("payload", [{}, {"period": 2026}, {"period": PERIOD, "surprise": "field"}])
def test_malformed_json_has_the_workbench_error_shape(client, payload):
    response = client.post(PREFIX + "/validate-saved", json=payload)
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "INVALID_REQUEST"


def test_ambiguous_dates_need_explicit_interpretation(client):
    files = monthly_files(sales=source([transaction_row(invoice_date="09/01/2026")]))
    review = validate(client, files)
    assert not review["ready"]
    assert any("AMBIGUOUS" in item["code"] for item in review["errors"])
    chosen = validate(client, files, **{"sales_history.date_format": "%m/%d/%Y"})
    assert chosen["ready"], chosen
    result = generate(client, chosen)
    stored = library.get_version("sales_history", result["committed_versions"]["sales_history"])
    assert stored.date_format == "%m/%d/%Y"
    replay = client.post(PREFIX + "/validate-saved", json={"period": PERIOD, "cycle_id": result["receipt"]["cycle_id"]}).json()
    assert preview(client, generate(client, replay)) == preview(client, result)


def test_cycle_receipt_is_immutable_and_rejects_corrupted_identity(client):
    _, result = cycle(client)
    from app.models.monthly_workflow import CycleReceipt
    receipt = CycleReceipt.model_validate(result["receipt"])
    with pytest.raises(DataLibraryError): cycle_receipts.CYCLE_RECEIPTS.save(receipt)
    path = cycle_receipts.CYCLE_RECEIPTS.root / PERIOD / (receipt.cycle_id + ".json")
    path.write_text(receipt.model_copy(update={"period": "2026-08"}).model_dump_json())
    with pytest.raises(DataLibraryError): cycle_receipts.CYCLE_RECEIPTS.get(PERIOD, receipt.cycle_id)


def test_a_saved_cycle_warns_when_the_report_action_version_changes(client, monkeypatch):
    _, result = cycle(client)
    from app.actions import registry
    from app.models.report_spec import REPORT_ACTION_ID
    installed = registry.get_action(REPORT_ACTION_ID)
    assert installed is not None
    monkeypatch.setattr(installed, "version", "0.4.0")
    review = client.post(PREFIX + "/validate-saved", json={"period": PERIOD, "cycle_id": result["receipt"]["cycle_id"]}).json()
    assert review["ready"]
    assert "ACTION_VERSION_CHANGED" in {item["code"] for item in review["warnings"]}
    assert review["action"]["version"] == "0.4.0"


def test_library_write_failure_does_not_disclose_a_physical_path(client, monkeypatch, data_library):
    review = validate(client, monthly_files())
    from pathlib import Path
    original_mkdir = Path.mkdir
    def denied(path, *args, **kwargs):
        if path == data_library.root:
            raise PermissionError("Cannot create /Users/secret/company/data/library")
        return original_mkdir(path, *args, **kwargs)
    monkeypatch.setattr(Path, "mkdir", denied)
    result = generate(client, review)
    assert result["status"] == "commit_failed"
    assert "secret" not in str(result)
    assert result["error"]["details"]["reason"] == "PermissionError"


def test_mixed_excel_types_are_disclosed_before_committing(client):
    rows = [transaction_row(invoice_date="2026-09-01", vintage=2021),
        transaction_row(invoice_date="2026-09-02", vintage="NV", invoice_number="second")]
    files = monthly_files()
    files["sales_history"] = ("mixed.xlsx", transactions(rows).as_xlsx())
    review = validate(client, files)
    assert review["ready"], review
    assert "MIXED_COLUMN_TYPES" in {issue["code"] for issue in review["warnings"]}
    assert library.list_datasets() == []


def test_history_discard_releases_the_parsed_review(client):
    response = client.post(PREFIX + "/history/validate", data={"dataset_id": "sales_history"},
        files={"source_file": ("history.csv", source(golden.SALES_ROWS["2025-08"]))})
    assert response.status_code == 200 and response.json()["ready"]
    token = response.json()["validation_id"]
    client.post(PREFIX + "/discard", json={"validation_id": token})
    assert history_workflow.HISTORY_WORKFLOW._pending is None
    assert client.post(PREFIX + "/history/commit", json={"validation_id": token}).status_code == 400
