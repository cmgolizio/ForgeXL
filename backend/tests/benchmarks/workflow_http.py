"""Verify production Next.js → loopback FastAPI → disk → XLSX/ZIP.

Run after npm run build, from backend:
    .venv/bin/python -m tests.benchmarks.workflow_http
Servers share this process's network namespace. An isolated temporary library
and synthetic golden fixtures protect all real business data.
"""
from __future__ import annotations

from io import BytesIO
import json
import os
from pathlib import Path
import socket
import shutil
import subprocess
import sys
from tempfile import TemporaryDirectory
import time
import zipfile

import httpx
from openpyxl import load_workbook

from tests.fixtures import report_months as golden
from tests.fixtures.monthly_sources import transactions
from tests.helpers import csv_bytes, xlsx_bytes


def verify_proof_actions(client):
    """Real production transport, both formats, independent literal controls."""
    headers = ["SKU", "Vintage", "Supplier", "Producer", "Selection", "Volume", "Invoice"]
    rows = [["A1", 2024, "Supplier", "Pród", "Cuvée", "750mL", "I1"],
        ["A1", 2024, "Supplier", "Pród", "Cuvée", "750mL", "I1"],
        ["A1", 2024, "Supplier", "Pród", "Cuvée", "750mL", "I2"]]
    for extension, payload in (("csv", csv_bytes(headers, rows)),
        ("xlsx", xlsx_bytes({"Sales": [headers, *rows]}))):
        for action, slot, output, expected in (
            ("exact_duplicate_remover", "source_file", "deduplicated_data", [rows[0], rows[2]]),
            ("product_master_builder", "sales_file", "product_master", [rows[0][:6]])):
            manifest = checked(client.post("/forge-api/api/runs", data={"action_id": action},
                files={slot: (f"proof.{extension}", payload)}))
            base = f'/forge-api/api/runs/{manifest["run_id"]}'
            page = checked(client.get(f"{base}/outputs/{output}/preview"))
            assert page["rows"] == expected
            csv = client.get(f"{base}/outputs/{output}/download/csv")
            assert csv.status_code == 200 and "Pród" in csv.text
            workbook_response = client.get(f"{base}/outputs/{output}/download/xlsx")
            assert workbook_response.status_code == 200
            workbook = load_workbook(BytesIO(workbook_response.content), data_only=True)
            sheet = workbook.active
            assert sheet is not None
            assert [list(row) for row in sheet.iter_rows(min_row=2, values_only=True)] == expected
            workbook.close()
            assert checked(client.post(base + "/discard"))["discarded"]
            assert client.get(base).status_code == 404
    invalid = client.post("/forge-api/api/runs", data={"action_id": "product_master_builder"},
        files={"sales_file": ("invalid.csv", b"SKU\nA1\n")})
    assert invalid.status_code == 422 and invalid.json()["error"]["code"] == "MISSING_COLUMNS"


def port():
    with socket.socket() as listener:
        listener.bind(("127.0.0.1", 0))
        return listener.getsockname()[1]


def checked(response):
    assert response.status_code == 200, response.text
    return response.json()


def unpack(response):
    assert response.status_code == 200 and "September 2026" in response.headers["content-disposition"]
    with zipfile.ZipFile(BytesIO(response.content)) as bundle:
        return {name: bundle.read(name) for name in bundle.namelist()}


def main():
    repository = Path(__file__).resolve().parents[3]
    api_port, web_port = port(), port()
    processes = []
    with TemporaryDirectory(prefix="forgexl-http-verification-") as temporary:
        log_path = Path(temporary) / "server.log"
        with log_path.open("w+") as log:
            environment = {**os.environ, "FORGEXL_LIBRARY_DIRECTORY": str(Path(temporary) / "library"),
                "FORGEXL_BACKEND_PORT": str(api_port), "NEXT_TELEMETRY_DISABLED": "1"}
            def start_api():
                process = subprocess.Popen([sys.executable, "-m", "uvicorn", "app.main:app", "--host", "127.0.0.1", "--port", str(api_port)],
                    cwd=repository / "backend", env=environment, stdout=log, stderr=log)
                processes.append(process)
                return process
            try:
                backend = start_api()
                processes.append(subprocess.Popen(["node", "node_modules/next/dist/bin/next", "start", "--hostname", "127.0.0.1", "--port", str(web_port)],
                    cwd=repository, env=environment, stdout=log, stderr=log))
                with httpx.Client(base_url=f"http://127.0.0.1:{web_port}", timeout=30, trust_env=False) as client:
                    def wait_ready():
                        for _ in range(80):
                            try:
                                if client.get("/forge-api/health").status_code == 200: return
                            except httpx.HTTPError: pass
                            time.sleep(.1)
                        log.flush()
                        raise AssertionError("Server did not become ready: " + log_path.read_text())
                    wait_ready()
                    page = client.get("/monthly-reports")
                    assert page.status_code == 200 and "Monthly sales rep reports" in page.text
                    assert "What would you like to create?" in client.get("/").text
                    verify_proof_actions(client)
                    base = "/forge-api/api/monthly"
                    denied = client.post(base + "/validate", headers={"Origin": "https://untrusted.example"}, data={"period": golden.GOLDEN_MONTH})
                    assert denied.status_code == 403 and denied.json()["error"]["code"] == "CROSS_ORIGIN_REQUEST"
                    assert not (Path(temporary) / "library").exists()
                    # Legitimate browser Origin is the web server's random port,
                    # not the backend allowlist: the proxy validates then strips it.
                    client.headers["Origin"] = f"http://127.0.0.1:{web_port}"
                    historical = sorted(month for month in golden.SALES_ROWS if month != golden.GOLDEN_MONTH)
                    split = len(historical) // 2
                    for dataset_id, rows, skip in (
                        ("sales_history", [row for month in historical[:split] for row in golden.SALES_ROWS[month]], False),
                        ("sales_history", [row for month in historical[split - 1:] for row in golden.SALES_ROWS[month]], True),
                        ("sample_history", list(golden.SAMPLE_ROWS["2026-08"]), False)):
                        if skip:
                            blocked = checked(client.post(base + "/history/validate", data={"dataset_id": dataset_id},
                                files={"source_file": ("history.csv", transactions(rows).as_csv())}))
                            assert not blocked["ready"] and blocked["validation_id"] is None
                        review = checked(client.post(base + "/history/validate", data={"dataset_id": dataset_id, "skip_existing": str(skip).lower()},
                            files={"source_file": ("history.csv", transactions(rows).as_csv())}))
                        assert review["ready"], review
                        if skip: assert review["skipped_periods"] == [historical[split - 1]]
                        saved = checked(client.post(base + "/history/commit", json={"validation_id": review["validation_id"], "acknowledge_warnings": True}))
                        assert saved["status"] == "saved"
                    review = checked(client.post(base + "/validate", data={"period": golden.GOLDEN_MONTH}, files={
                        "sales_history": ("sales.csv", golden.sales_table(golden.GOLDEN_MONTH).as_csv()),
                        "sample_history": ("samples.csv", golden.sample_table(golden.GOLDEN_MONTH).as_csv())}))
                    assert review["ready"], review
                    result = checked(client.post(base + "/generate", json={"validation_id": review["validation_id"], "acknowledge_warnings": True}))
                    assert result["status"] == "reports_generated"
                    run_id = result["manifest"]["run_id"]
                    original = unpack(client.get(f"/forge-api/api/runs/{run_id}/artifacts/download/zip"))
                    assert len(original) == 2
                    for payload in original.values():
                        workbook = load_workbook(BytesIO(payload), data_only=True)
                        assert len(workbook.sheetnames) == 6
                        workbook.close()
                    summary = checked(client.get(f"/forge-api/api/runs/{run_id}/outputs/company_summary/preview"))
                    assert summary["rows"][0][summary["columns"].index("Revenue")] == 995.0
                    assert checked(client.post(f"/forge-api/api/runs/{run_id}/discard"))["discarded"]
                    assert client.get(f"/forge-api/api/runs/{run_id}").status_code == 404
                    assert checked(client.get(base + "/catalog"))["periods"]
                    backend.terminate(); backend.wait(timeout=5)
                    unavailable = client.get("/forge-api/health")
                    assert unavailable.status_code == 502
                    # Restore a stopped full-library backup, including receipts.
                    restored = Path(temporary) / "restored-library"
                    shutil.copytree(Path(temporary) / "library", restored)
                    environment["FORGEXL_LIBRARY_DIRECTORY"] = str(restored)
                    start_api(); wait_ready()
                    assert client.get(f"/forge-api/api/runs/{run_id}").status_code == 404
                    replay = checked(client.post(base + "/validate-saved", json={"period": golden.GOLDEN_MONTH, "cycle_id": result["receipt"]["cycle_id"]}))
                    repeated = checked(client.post(base + "/generate", json={"validation_id": replay["validation_id"], "acknowledge_warnings": True}))
                    assert repeated["status"] == "reports_generated"
                    assert unpack(client.get(f'/forge-api/api/runs/{repeated["manifest"]["run_id"]}/artifacts/download/zip')) == original
                    print(json.dumps({"http_proxy": "passed", "proof_actions_csv_xlsx": "passed",
                        "cross_origin_write_refused": "passed", "same_origin_browser_write": "passed",
                        "disconnected_backend_502": "passed", "explicit_run_release_preserves_saved_cycle": "passed",
                        "history_chunks_overlap_consent": "passed", "monthly_cycle": "passed",
                        "workbooks": 2, "worksheets": 12, "company_revenue_control": 995.0,
                        "restart_and_backup_restore_exact_workbook_replay": "passed", "live_browser": "not exercised"}, indent=2))
            finally:
                for process in processes:
                    if process.poll() is None:
                        process.terminate()
                        try: process.wait(timeout=5)
                        except subprocess.TimeoutExpired: process.kill(); process.wait(timeout=5)


if __name__ == "__main__": main()
